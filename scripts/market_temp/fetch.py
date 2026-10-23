#!/usr/bin/env python
"""Fetch layer for the market-temperature pipeline.

Every source is wrapped in try/except. On success the series is rewritten in data/market_temp/;
on failure the previous cache is kept (or seeded from the earlier backtests/ folders on first run)
and the manifest records the error. compute.py then works from whatever is in the cache, and
needs_ai.json lists the items that are failed or stale so the AI step knows what to go and find.

Sources (same endpoints as the 2026-10 studies, so numbers reconcile):
  Yahoo v8 chart        ^GSPC ^SP500TR SPY ^NDX QQQ ^W5000 ^TNX ^IRX ^FVX ^VIX GC=F DX-Y.NYB   (live in-progress row dropped)
  BLS API v1            CPI-U NSA CUUR0000SA0, unemployment SA LNS14000000 (10-year chunks, no key)
  historyofmarket.com   /api/ndx/forward-pe.json  (NDX forward PE weekly since 2001 + trailing PE daily; CC BY 4.0)
  multpl.com            Shiller CAPE and S&P 500 trailing PE, monthly tables
  FRED                  M2SL via fredgraph.csv (usually unreachable from this machine -> cache + AI flag)
  BEA                   gdplev.xlsx (annual + quarterly nominal/real GDP)
  LBMA                  gold PM fix json (403 from here -> cache; compute.py splices COMEX GC=F after the cache ends)
  stockanalysis.com     ETF holdings tables for QQQ and SPY (top weights -> concentration)
"""
import csv
import datetime as dt
import io
import json
import re
import shutil
import sys
import urllib.parse
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import DATA, REPO, http_get, log, read_json, utc_now, write_json  # noqa: E402

MANIFEST = DATA / "manifest.json"

YAHOO = {"GSPC": "^GSPC", "SP500TR": "^SP500TR", "SPY": "SPY", "NDX": "^NDX", "QQQ": "QQQ", "W5000": "^W5000",
         "TNX": "^TNX", "IRX": "^IRX", "FVX": "^FVX", "VIX": "^VIX", "GCF": "GC=F", "DXY": "DX-Y.NYB"}
YAHOO_TPL = ("https://query1.finance.yahoo.com/v8/finance/chart/{s}?period1=0&period2={p2}"
             "&interval=1d&events=div%2Csplits&includeAdjustedClose=true")

# first-run seeds: copied from the earlier studies when the live fetch fails and no cache exists yet
SEEDS = {
    "us_m2_monthly.csv": REPO / "backtests" / "黄金_货币供应_估值_2026-08-07" / "data" / "us_m2_monthly.csv",
    "gold_lbma_pm_daily.csv": REPO / "backtests" / "黄金_货币供应_估值_2026-08-07" / "data" / "gold_lbma_pm_daily.csv",
    "bea_gdp_quarterly.csv": REPO / "backtests" / "QQQ_vs_GDP加CPI_锚点2010_2020_2026-10-12" / "data" / "bea_gdp_quarterly.csv",
    "bea_gdp_annual.csv": REPO / "backtests" / "QQQ_vs_GDP加CPI_锚点2010_2020_2026-10-12" / "data" / "bea_gdp_annual.csv",
    "qqq_holdings_top.csv": REPO / "backtests" / "巨头市值占比_历史对照_2026-10-12" / "data" / "qqq_holdings_top.csv",
    "spy_holdings_top.csv": REPO / "backtests" / "巨头市值占比_历史对照_2026-10-12" / "data" / "spy_holdings_top.csv",
    "bls_cpi_unrate_monthly.csv": REPO / "backtests" / "市场体温_三节点_1980起_2026-10-11" / "data" / "bls_cpi_unrate_monthly.csv",
    "multpl_shiller_pe.csv": REPO / "backtests" / "市场体温_三节点_1980起_2026-10-11" / "data" / "multpl_shiller_pe.csv",
    "multpl_sp500_pe.csv": REPO / "backtests" / "市场体温_三节点_1980起_2026-10-11" / "data" / "multpl_sp500_pe.csv",
    "ndx_forward_pe_weekly.csv": REPO / "backtests" / "市场体温_三节点_1980起_2026-10-11" / "data" / "ndx_forward_pe_weekly.csv",
}

# staleness thresholds (days) beyond which an item goes on the AI list even if the cache exists
STALE_DAYS = {"yahoo": 5, "bls": 60, "hom_forward_pe": 14, "multpl_cape": 45, "multpl_pe": 45,
              "fred_m2": 60, "bea_gdp": 120, "lbma_gold": 10, "sa_qqq": 30, "sa_spy": 30}


class Manifest:
    def __init__(self):
        self.m = read_json(MANIFEST, {"series": {}})
        self.m["fetched_utc"] = utc_now()

    def ok(self, key, **kw):
        rec = dict(status="ok", fetched_utc=utc_now(), error=None, **kw)
        self.m["series"][key] = rec
        log("OK  ", key, {k: v for k, v in kw.items() if k in ("rows", "last", "url")})

    def fail(self, key, err, **kw):
        prev = self.m["series"].get(key, {})
        rec = dict(prev)
        rec.update(status="cached" if prev.get("last") else "failed", error=str(err)[:300], failed_utc=utc_now(), **kw)
        self.m["series"][key] = rec
        log("FAIL", key, str(err)[:160], "-> using cache" if prev.get("last") else "-> NO CACHE")

    def save(self):
        write_json(MANIFEST, self.m)


def _seed(fname, man: Manifest, key):
    """copy a seed file into the cache when the live fetch failed and nothing is cached yet"""
    dst = DATA / fname
    if dst.exists():
        return
    src = SEEDS.get(fname)
    if src and src.exists():
        shutil.copy(src, dst)
        rec = man.m["series"].setdefault(key, {})
        rec.update(status="seeded", seeded_from=str(src.relative_to(REPO)))
        try:
            last = _last_date_in_csv(dst)
            rec["last"] = last
        except Exception:
            pass
        log("SEED", key, "copied", src.relative_to(REPO))


def _last_date_in_csv(path: Path) -> str:
    rows = list(csv.reader(path.open(encoding="utf-8-sig")))
    return rows[-1][0] if len(rows) > 1 else ""


# ---------------------------------------------------------------- Yahoo
def fetch_yahoo(man: Manifest):
    p2 = int((dt.datetime.utcnow() + dt.timedelta(days=2)).timestamp())
    for label, sym in YAHOO.items():
        url = YAHOO_TPL.format(s=urllib.parse.quote(sym), p2=p2)
        try:
            js = json.loads(http_get(url, timeout=60).decode())
            res = js["chart"]["result"][0]
            ts = res["timestamp"]
            q = res["indicators"]["quote"][0]["close"]
            adj = res["indicators"].get("adjclose", [{}])[0].get("adjclose", q)
            meta = res["meta"]
            reg = meta.get("currentTradingPeriod", {}).get("regular", {})
            live = meta.get("regularMarketTime", 0) < reg.get("end", 0) and len(ts) and ts[-1] >= reg.get("start", 10**12)
            if live:
                ts, q, adj = ts[:-1], q[:-1], adj[:-1]
                log("     ", label, "dropped live in-progress row")
            rows = []
            for t, c, a in zip(ts, q, adj):
                if c is None:
                    continue
                d = dt.datetime.utcfromtimestamp(t).date().isoformat()
                rows.append((d, float(c), float(a if a is not None else c)))
            if len(rows) < 100:
                raise ValueError(f"only {len(rows)} rows")
            with (DATA / f"{label}_daily.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["date", "close", "adj_close"])
                w.writerows(rows)
            man.ok(f"yahoo_{label}", symbol=sym, url=url, rows=len(rows), first=rows[0][0], last=rows[-1][0],
                   last_close=rows[-1][1], kind="yahoo")
        except Exception as e:  # noqa: BLE001
            man.fail(f"yahoo_{label}", e, symbol=sym, url=url, kind="yahoo")


# ---------------------------------------------------------------- BLS
def fetch_bls(man: Manifest):
    url = "https://api.bls.gov/publicAPI/v1/timeseries/data/"
    series = {"CUUR0000SA0": "cpi_u_nsa", "LNS14000000": "unrate_sa"}
    try:
        out = {}
        year, end_year = 1978, dt.date.today().year
        while year <= end_year:
            chunk_end = min(year + 9, end_year)
            payload = json.dumps({"seriesid": list(series), "startyear": str(year), "endyear": str(chunk_end)}).encode()
            body = json.loads(http_get(url, data=payload, headers={"Content-Type": "application/json"}, timeout=90).decode())
            if body.get("status") != "REQUEST_SUCCEEDED":
                raise RuntimeError(f"BLS {year}-{chunk_end}: {body.get('message')}")
            for s in body["Results"]["series"]:
                col = series[s["seriesID"]]
                for o in s["data"]:
                    if not o["period"].startswith("M") or o["period"] == "M13":
                        continue
                    try:
                        out.setdefault(f"{o['year']}-{o['period'][1:]}", {})[col] = float(o["value"])
                    except ValueError:
                        continue
            year = chunk_end + 1
        rows = sorted(out.items())
        if len(rows) < 400:
            raise ValueError(f"only {len(rows)} months")
        with (DATA / "bls_cpi_unrate_monthly.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["month", "cpi_u_nsa", "unrate_sa"])
            for m, d in rows:
                w.writerow([m, d.get("cpi_u_nsa", ""), d.get("unrate_sa", "")])
        man.ok("bls", url=url, rows=len(rows), first=rows[0][0], last=rows[-1][0], kind="bls",
               last_values=rows[-1][1])
    except Exception as e:  # noqa: BLE001
        man.fail("bls", e, url=url, kind="bls")
        _seed("bls_cpi_unrate_monthly.csv", man, "bls")


# ---------------------------------------------------------------- historyofmarket NDX PE
def fetch_hom(man: Manifest):
    url = "https://historyofmarket.com/api/ndx/forward-pe.json"
    try:
        j = json.loads(http_get(url, timeout=60).decode())
        fwd = sorted((x["date"], float(x["value"])) for x in j["forward"] if x.get("value") is not None)
        trl = sorted((x["date"], float(x["value"])) for x in j.get("trailing", []) if x.get("value") is not None)
        if len(fwd) < 500:
            raise ValueError(f"forward list too short: {len(fwd)}")
        # the weekly consensus series lags; historyofmarket's own daily blended FY1/FY2 estimate continues it (same as the 10-08 study)
        own = sorted((x["date"], float(x["value"])) for x in j.get("forwardOwn", []) if x.get("value") is not None and x["date"] > fwd[-1][0])
        splice = dict(weekly_last=fwd[-1][0], weekly_last_value=fwd[-1][1], own_from=own[0][0] if own else None, own_first_value=own[0][1] if own else None,
                      own_last=own[-1][0] if own else None, own_basis=(j.get("forwardOwn") or [{}])[-1].get("basis"))
        fwd = fwd + own
        for name, rows in (("ndx_forward_pe_weekly.csv", fwd), ("ndx_trailing_pe_daily.csv", trl)):
            with (DATA / name).open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["date", "value"])
                w.writerows(rows)
        (DATA / "hom_ndx_forward_pe.json").write_text(json.dumps(j, ensure_ascii=False), encoding="utf-8")
        man.ok("hom_forward_pe", url=url, rows=len(fwd), first=fwd[0][0], last=fwd[-1][0], kind="hom_forward_pe",
               updated=j.get("updated"), current=j.get("current"), splice=splice, attribution=j.get("_attribution"), license="CC BY 4.0")
    except Exception as e:  # noqa: BLE001
        man.fail("hom_forward_pe", e, url=url, kind="hom_forward_pe")
        _seed("ndx_forward_pe_weekly.csv", man, "hom_forward_pe")


# ---------------------------------------------------------------- multpl
def _parse_multpl(html: str):
    pat = re.compile(r"<td>([A-Z][a-z]{2} \d{1,2}, \d{4})</td>\s*<td>\s*(?:<abbr[^>]*>[^<]*</abbr>|&#x2002;|&nbsp;)?\s*([\d.]+)", re.S)
    rows = []
    for d, v in pat.findall(html):
        try:
            rows.append((dt.datetime.strptime(d, "%b %d, %Y").date().isoformat(), float(v)))
        except ValueError:
            continue
    return sorted(set(rows))


def fetch_multpl(man: Manifest):
    for key, url, fname in (("multpl_cape", "https://www.multpl.com/shiller-pe/table/by-month", "multpl_shiller_pe.csv"),
                            ("multpl_pe", "https://www.multpl.com/s-p-500-pe-ratio/table/by-month", "multpl_sp500_pe.csv")):
        try:
            rows = _parse_multpl(http_get(url, timeout=60).decode("utf-8", "ignore"))
            if len(rows) < 1000:
                raise ValueError(f"parsed only {len(rows)} rows")
            with (DATA / fname).open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["date", "value"])
                w.writerows(rows)
            man.ok(key, url=url, rows=len(rows), first=rows[0][0], last=rows[-1][0], last_value=rows[-1][1], kind=key)
        except Exception as e:  # noqa: BLE001
            man.fail(key, e, url=url, kind=key)
            _seed(fname, man, key)


# ---------------------------------------------------------------- FRED M2
def fetch_fred_m2(man: Manifest):
    url = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=M2SL"
    try:
        txt = http_get(url, timeout=15, retries=0).decode()
        rows = []
        for line in txt.splitlines()[1:]:
            d, v = line.split(",")[:2]
            if v in ("", "."):
                continue
            rows.append((d[:7], float(v)))
        if len(rows) < 600:
            raise ValueError(f"only {len(rows)} rows")
        with (DATA / "us_m2_monthly.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["month", "m2_sa_usd_bn"])
            w.writerows(rows)
        man.ok("fred_m2", url=url, rows=len(rows), first=rows[0][0], last=rows[-1][0], kind="fred_m2")
    except Exception as e:  # noqa: BLE001
        man.fail("fred_m2", e, url=url, kind="fred_m2")
        _seed("us_m2_monthly.csv", man, "fred_m2")


# ---------------------------------------------------------------- BEA GDP
NIPA_Q = "https://apps.bea.gov/national/Release/TXT/NipaDataQ.txt"   # ~35 MB flat file, monthly; A191RC nominal / A191RX real GDP, millions SAAR


def _qend(per):
    y, n = int(per[:4]), int(per[-1])
    month = n * 3
    return (dt.date(y + (month == 12), (month % 12) + 1, 1) - dt.timedelta(days=1)).isoformat()


def _parse_nipa_q(txt):
    rows = {}
    for line in txt.splitlines():
        if line.startswith("A191RC,") or line.startswith("A191RX,"):
            code, per, val = line.split(",", 2)
            try:
                rows.setdefault(per, {})[code] = float(val.strip().strip('"').replace(",", "")) / 1000.0
            except ValueError:
                continue
    out = []
    for per in sorted(rows):
        r = rows[per]
        if "A191RC" in r and "A191RX" in r and re.fullmatch(r"\d{4}Q[1-4]", per):
            out.append((per.lower(), r["A191RC"], r["A191RX"], _qend(per)))
    return out


def _annual_from_quarterly(qtr):
    by = {}
    for q, nom, real, _d in qtr:
        by.setdefault(int(q[:4]), []).append((nom, real))
    return [(y, sum(v[0] for v in vs) / 4, sum(v[1] for v in vs) / 4) for y, vs in sorted(by.items()) if len(vs) == 4]


def fetch_bea(man: Manifest):
    url = "https://apps.bea.gov/national/xls/gdplev.xlsx"
    src = url
    try:
        import openpyxl
        raw = http_get(url, timeout=60)
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        ann, qtr = [], []
        for row in ws.iter_rows(values_only=True):
            if row[0] and re.fullmatch(r"\d{4}", str(row[0]).strip()) and row[1] and row[2]:
                ann.append((int(str(row[0]).strip()), float(row[1]), float(row[2])))
            if len(row) > 6 and row[4] and re.fullmatch(r"\d{4}[Qq][1-4]", str(row[4]).strip()) and row[5] and row[6]:
                q = str(row[4]).strip().upper()
                y, n = int(q[:4]), int(q[-1])
                month = n * 3
                last_day = (dt.date(y + (month == 12), (month % 12) + 1, 1) - dt.timedelta(days=1))
                qtr.append((q.lower(), float(row[5]), float(row[6]), last_day.isoformat()))
        if len(qtr) < 300 or len(ann) < 90:
            raise ValueError(f"parsed ann={len(ann)} qtr={len(qtr)}")
        # gdplev.xlsx has not been refreshed since 2025Q1; when it is >150 days stale, read the quarterly NIPA flat file instead
        if (dt.date.today() - dt.date.fromisoformat(qtr[-1][3])).days > 150:
            try:
                q2 = _parse_nipa_q(http_get(NIPA_Q, timeout=240, retries=0).decode("utf-8", "ignore"))
                if len(q2) > len(qtr):
                    log("     ", "bea_gdp: gdplev stale, using NipaDataQ.txt through", q2[-1][0])
                    qtr, ann, src = q2, _annual_from_quarterly(q2), NIPA_Q
            except Exception as e2:  # noqa: BLE001
                log("     ", "bea_gdp: NipaDataQ fallback failed:", str(e2)[:120])
        with (DATA / "bea_gdp_quarterly.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["quarter", "gdp_nominal_bn", "gdp_real_bn", "date"])
            w.writerows(qtr)
        with (DATA / "bea_gdp_annual.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["year", "gdp_nominal_bn", "gdp_real_bn"])
            w.writerows(ann)
        (DATA / "gdplev.xlsx").write_bytes(raw)
        man.ok("bea_gdp", url=src, rows=len(qtr), first=qtr[0][0], last=qtr[-1][3], last_quarter=qtr[-1][0],
               last_nominal_bn=qtr[-1][1], kind="bea_gdp", annual_rows=len(ann))
    except Exception as e:  # noqa: BLE001
        man.fail("bea_gdp", e, url=url, kind="bea_gdp")
        _seed("bea_gdp_quarterly.csv", man, "bea_gdp")
        _seed("bea_gdp_annual.csv", man, "bea_gdp")


# ---------------------------------------------------------------- LBMA gold
def fetch_lbma(man: Manifest):
    url = "https://prices.lbma.org.uk/json/gold_pm.json"
    try:
        j = json.loads(http_get(url, timeout=60, retries=0).decode())
        g = sorted((x["d"], float(x["v"][0])) for x in j if x.get("v") and x["v"][0])
        if len(g) < 5000:
            raise ValueError(f"only {len(g)} rows")
        with (DATA / "gold_lbma_pm_daily.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["date", "usd_per_oz"])
            w.writerows(g)
        man.ok("lbma_gold", url=url, rows=len(g), first=g[0][0], last=g[-1][0], kind="lbma_gold")
    except Exception as e:  # noqa: BLE001
        man.fail("lbma_gold", e, url=url, kind="lbma_gold", note="compute.py splices COMEX GC=F after the cache ends")
        _seed("gold_lbma_pm_daily.csv", man, "lbma_gold")


# ---------------------------------------------------------------- stockanalysis ETF holdings
def _parse_holdings(html: str):
    pat = re.compile(r'<a href="/stocks/([a-z0-9.\-]+)/"\s*>([A-Z0-9.\-]{1,6})</a>(?:(?!</tr>).)*?<td class="[a-z]+ svelte-[a-z0-9]+">([^<]*)</td>(?:(?!</tr>).)*?<td class="svelte-[a-z0-9]+">([\d.]+)%', re.S)
    rows, seen = [], set()
    for _slug, sym, name, w in pat.findall(html):
        if sym in seen:
            continue
        seen.add(sym)
        rows.append((sym, float(w), name.strip()))
    return rows


def fetch_stockanalysis(man: Manifest):
    for etf in ("qqq", "spy"):
        url = f"https://stockanalysis.com/etf/{etf}/holdings/"
        key = f"sa_{etf}"
        try:
            rows = _parse_holdings(http_get(url, timeout=60, headers={"Referer": "https://stockanalysis.com/"}).decode("utf-8", "ignore"))
            if len(rows) < 10 or rows[0][1] < 3:
                raise ValueError(f"parsed {len(rows)} rows, top={rows[:2]}")
            with (DATA / f"{etf}_holdings_top.csv").open("w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(["symbol", "weight", "name"])
                w.writerows(rows)
            man.ok(key, url=url, rows=len(rows), last=dt.date.today().isoformat(), top=rows[:3], kind=key)
        except Exception as e:  # noqa: BLE001
            man.fail(key, e, url=url, kind=key)
            _seed(f"{etf}_holdings_top.csv", man, key)


# ---------------------------------------------------------------- staleness -> needs_ai
def staleness(man: Manifest, as_of: str):
    """mark each series stale/fresh relative to as_of and build the AI to-do list"""
    ref = dt.date.fromisoformat(as_of)
    needs = []
    for key, rec in man.m["series"].items():
        kind = rec.get("kind", key)
        last = rec.get("last")
        thr = STALE_DAYS.get(kind, 30)
        days = None
        if last:
            try:
                ls = str(last)[:10]
                if len(ls) == 7:  # BLS / M2 months come as YYYY-MM
                    ls += "-01"
                days = (ref - dt.date.fromisoformat(ls)).days
            except ValueError:
                days = None
        rec["stale_days"] = days
        rec["stale"] = (days is None) or (days > thr)
        if rec.get("status") in ("failed", "cached", "seeded") or rec["stale"]:
            needs.append(dict(key=key, status=rec.get("status"), last=last, stale_days=days, threshold_days=thr,
                              url=rec.get("url"), error=rec.get("error"), what=_what(kind)))
    man.m["needs_ai"] = needs
    return needs


def _what(kind):
    return {
        "fred_m2": "美国 M2（季调，十亿美元，月频）。FRED 本机不可达：请从 federalreserve.gov H.6 或其它镜像取最近几个月的 M2SL，写明来源与发布日期。",
        "lbma_gold": "LBMA 黄金 PM 定盘价。本机 403：pipeline 已用 COMEX GC=F 接续，只需确认接续日前后价差 <1%。",
        "bea_gdp": "BEA 名义/实际 GDP 季度值。若新季度已发布而缓存没有，补最新一季（SAAR，十亿美元）。",
        "bls": "BLS CPI-U 与失业率最新月份。",
        "hom_forward_pe": "纳指100 前瞻 PE（historyofmarket.com）。若接口失效，用 FactSet Earnings Insight 的标普前瞻 PE 作旁证。",
        "yahoo": "Yahoo 行情。若失败，用其它行情源核对最近收盘。",
        "sa_qqq": "QQQ 前十大持仓权重（stockanalysis 或 Invesco 官网）。",
        "sa_spy": "SPY 前十大持仓权重（stockanalysis 或 SSGA 官网）。",
        "multpl_cape": "标普500 CAPE（multpl.com 或 Shiller 数据）。",
        "multpl_pe": "标普500 滚动 PE（multpl.com）。",
    }.get(kind, "核对并补齐该数据。")


def fetch_all(skip_network=False):
    man = Manifest()
    if skip_network:
        log("fetch skipped (--no-fetch); using cache in", DATA)
        for fname in SEEDS:
            _seed(fname, man, fname.split(".")[0])
    else:
        fetch_yahoo(man)
        fetch_bls(man)
        fetch_hom(man)
        fetch_multpl(man)
        fetch_fred_m2(man)
        fetch_bea(man)
        fetch_lbma(man)
        fetch_stockanalysis(man)
    man.save()
    return man


if __name__ == "__main__":
    m = fetch_all("--no-fetch" in sys.argv)
    print(json.dumps({k: v.get("status") for k, v in m.m["series"].items()}, indent=1))
